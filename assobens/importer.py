"""AssobensImporter: UPSERT por CHASSI na matriz publicada, com publicação atômica.

- chave única: CHASSI. Mesmo arquivo importado duas vezes => 0 inseridos, 0 atualizados.
- valor alterado no ASSOBENS durante o dia => registro ATUALIZADO, nunca duplicado.
- histórico anterior válido nunca é apagado: só se inserem/atualizam linhas.
- publicação: escreve em arquivo temporário no mesmo diretório e troca com os.replace;
  a matriz anterior é copiada para storage/app/assobens/backup antes da troca.
"""
from __future__ import annotations

import os
import shutil
import tempfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from . import config
from .errors import ImportError_

KEY = "CHASSI"


@dataclass
class UpsertResult:
    inserted: int = 0
    updated: int = 0
    unchanged: int = 0
    headers: list[str] = field(default_factory=list)
    rows: list[dict] = field(default_factory=list)   # matriz completa após o upsert

    @property
    def total(self) -> int:
        return len(self.rows)


class AssobensImporter:
    def __init__(self, matrix_path: Path | None = None, backup_dir: Path | None = None):
        self.matrix_path = Path(matrix_path or config.MATRIZ_PATH)
        self.backup_dir = Path(backup_dir or config.BACKUP_DIR)

    # ------------------------------------------------------------ leitura
    def load_matrix(self) -> tuple[list[str], list[dict]]:
        """Lê a matriz atual (cabeçalho + linhas como dict). Matriz inexistente = base vazia."""
        if not self.matrix_path.exists():
            return list(config.CANONICAL_COLUMNS), []
        import openpyxl
        wb = openpyxl.load_workbook(self.matrix_path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        headers = [str(h if h is not None else "").strip() for h in next(it)]
        rows = []
        for r in it:
            if r is None or not any(c not in (None, "") for c in r):
                continue
            d = {h: ("" if v is None else (v.isoformat() if isinstance(v, datetime) else str(v))) for h, v in zip(headers, r)}
            if d.get(KEY):
                rows.append(d)
        wb.close()
        return headers, rows

    # ------------------------------------------------------------ upsert
    def upsert(self, new_rows: list[dict], existing_headers: list[str] | None = None,
               existing_rows: list[dict] | None = None) -> UpsertResult:
        if existing_headers is None or existing_rows is None:
            existing_headers, existing_rows = self.load_matrix()
        headers = list(existing_headers) or list(config.CANONICAL_COLUMNS)
        for c in config.CANONICAL_COLUMNS:
            if c not in headers:
                headers.append(c)
        for r in new_rows:
            for c in r:
                if c not in headers:
                    headers.append(c)  # dimensão nova preservada
        index = {r[KEY]: i for i, r in enumerate(existing_rows)}
        merged = [dict(r) for r in existing_rows]
        res = UpsertResult(headers=headers)
        for r in new_rows:
            key = r.get(KEY)
            if not key:
                continue
            i = index.get(key)
            if i is None:
                merged.append({h: r.get(h, "") for h in headers})
                index[key] = len(merged) - 1
                res.inserted += 1
                continue
            cur = merged[i]
            changed = False
            for h in headers:
                nv = r.get(h)
                if nv is None or h not in r:
                    continue
                if not _same(cur.get(h, ""), nv):
                    cur[h] = nv
                    changed = True
            if changed:
                res.updated += 1
            else:
                res.unchanged += 1
        merged.sort(key=lambda x: (str(x.get("DATA EMPLACAMENTO", "")), str(x.get(KEY, ""))))
        res.rows = merged
        return res

    # ------------------------------------------------------------ publicação atômica
    def publish(self, headers: list[str], rows: list[dict], run_id: str = "") -> Path:
        if not rows:
            raise ImportError_("recusado publicar matriz vazia")
        import openpyxl
        self.matrix_path.parent.mkdir(parents=True, exist_ok=True)
        fd, tmp = tempfile.mkstemp(prefix=self.matrix_path.stem + ".", suffix=".tmp.xlsx", dir=self.matrix_path.parent)
        os.close(fd)
        try:
            wb = openpyxl.Workbook(write_only=True)
            ws = wb.create_sheet("Enriquecimento")
            ws.append(headers)
            for r in rows:
                ws.append([_cell(r.get(h, "")) for h in headers])
            wb.save(tmp)
            self._backup(run_id)
            os.replace(tmp, self.matrix_path)
        except Exception as e:
            try:
                os.unlink(tmp)
            except OSError:
                pass
            raise ImportError_(f"falha ao publicar matriz (a anterior permanece ativa): {e}") from e
        return self.matrix_path

    def _backup(self, run_id: str) -> None:
        if not self.matrix_path.exists():
            return
        self.backup_dir.mkdir(parents=True, exist_ok=True)
        stamp = run_id or datetime.now(config.TZ).strftime("%Y%m%d-%H%M%S")
        dst = self.backup_dir / f"{stamp}_{self.matrix_path.name}"
        shutil.copy2(self.matrix_path, dst)
        # retenção: mantém os 10 backups mais recentes
        olds = sorted(self.backup_dir.glob(f"*_{self.matrix_path.name}"))[:-10]
        for o in olds:
            try:
                o.unlink()
            except OSError:
                pass


def _cell(v):
    if v is None:
        return ""
    return str(v)


def _same(a, b) -> bool:
    """Igualdade que ignora espaços ocultos (NBSP, tabs, duplos): diferença só de espaçamento não é atualização."""
    norm = lambda v: " ".join(str("" if v is None else v).replace(" ", " ").split())
    return norm(a) == norm(b)
